"""
md_to_template_docx.py — Markdown 转 Word（保留模板格式）

解析 Markdown 文件，按标题层级映射到 Word 模板的章节结构，
将内容填入模板对应位置，保留模板全部格式。

用法：
    python md_to_template_docx.py --template 模板.docx --md 内容.md --output 输出.docx
    python md_to_template_docx.py --template 模板.docx --md 内容.md --mapping 映射.json

映射文件格式见 prompts/mapping_format.md。
"""
import argparse
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx_template import (
    load_template, save_doc, find_para_index, find_para_indices,
    clear_runs, copy_run_format, copy_para_format,
    add_styled_para, insert_paras_after, replace_para_text,
    add_image_after, get_normal_ref
)


def parse_markdown(md_path):
    """
    解析 Markdown 文件，返回结构化章节列表。
    每个章节: {'level': int, 'title': str, 'content': [str], 'images': [(alt, path)]}
    """
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    sections = []
    current = None

    for line in lines:
        line = line.rstrip('\n')

        # 标题匹配
        m = re.match(r'^(#{1,6})\s+(.+)$', line)
        if m:
            if current:
                sections.append(current)
            level = len(m.group(1))
            title = m.group(2).strip()
            current = {'level': level, 'title': title, 'content': [], 'images': []}
            continue

        if current is None:
            current = {'level': 0, 'title': '', 'content': [], 'images': []}

        # 图片匹配
        img_m = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line)
        if img_m:
            current['images'].append((img_m.group(1), img_m.group(2)))
            continue

        # HTML 注释中的图片引用 <!-- ![alt](path) -->
        comment_m = re.match(r'<!--\s*!\[([^\]]*)\]\(([^)]+)\)\s*-->', line)
        if comment_m:
            current['images'].append((comment_m.group(1), comment_m.group(2)))
            continue

        current['content'].append(line)

    if current:
        sections.append(current)

    return sections


def find_section_anchor(doc, section_title):
    """
    在模板中查找与 Markdown 章节标题匹配的锚点段落。
    尝试多种匹配策略：完全匹配、包含匹配、关键词匹配。
    """
    # 策略1：完全匹配
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == section_title:
            return i

    # 策略2：包含匹配
    for i, p in enumerate(doc.paragraphs):
        if section_title in p.text:
            return i

    # 策略3：关键词匹配（去掉标点和空格）
    clean_title = re.sub(r'[^\w]', '', section_title)
    for i, p in enumerate(doc.paragraphs):
        clean_text = re.sub(r'[^\w]', '', p.text)
        if clean_title and clean_title in clean_text:
            return i

    return None


def fill_section(doc, anchor_idx, paragraphs, ref_para=None, insert_after=True):
    """在锚点位置填充内容段落"""
    if ref_para is None:
        ref_para = doc.paragraphs[anchor_idx] if anchor_idx is not None else get_normal_ref(doc)

    if anchor_idx is None:
        # 没有找到锚点，追加到文档末尾
        for text in paragraphs:
            if text.strip():
                add_styled_para(doc, text, ref_para)
        return

    if insert_after:
        insert_paras_after(doc, anchor_idx, paragraphs, ref_para)
    else:
        # 替换锚点段落内容
        if paragraphs:
            replace_para_text(doc, anchor_idx, paragraphs[0])
            if len(paragraphs) > 1:
                insert_paras_after(doc, anchor_idx, paragraphs[1:], ref_para)


def fill_images(doc, anchor_idx, images, fig_dir=None):
    """在锚点位置插入图片"""
    if anchor_idx is None:
        return

    current_idx = anchor_idx
    for alt, path in images:
        # 解析相对路径
        if fig_dir and not os.path.isabs(path):
            path = os.path.join(fig_dir, path)
        if os.path.exists(path):
            img_para = add_image_after(doc, current_idx, path)
            current_idx += 1
            # 如果有 alt 文字（如"图1"），在图片前插入标签
            if alt and re.match(r'图\s*\d+', alt):
                from docx.shared import Pt
                label_para = doc.add_paragraph()
                label_para.alignment = __import__('docx.enum.text', fromlist=['WD_ALIGN_PARAGRAPH']).WD_ALIGN_PARAGRAPH.CENTER
                label_run = label_para.add_run(alt)
                label_run.font.bold = True
                # 移动到图片前
                ref_element = img_para._element
                parent = ref_element.getparent()
                parent.insert(list(parent).index(ref_element), label_para._element)
                current_idx += 1


# 默认映射：Markdown 标题 -> 模板中的锚点文本
DEFAULT_MAPPING = {
    '说明书摘要': '本发明公开了一种',
    '权利要求书': '1. 一种',
    '发明名称': '发明名称：',
    '技术领域': '本发明属于',
    '背景技术': '背景技术',
    '发明内容': '本发明的目的在于',
    '具体技术方案': '本发明所采用的具体技术方案如下',
    '有益效果': '本发明相对于现有技术而言',
    '附图说明': '图1为',
    '具体实施方式': '如图1所示，',
    '说明书附图': '说 明 书 附 图',
}


def load_mapping(mapping_path):
    """加载自定义映射文件"""
    if mapping_path and os.path.exists(mapping_path):
        with open(mapping_path, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}


def main():
    parser = argparse.ArgumentParser(description='Markdown 转 Word（保留模板格式）')
    parser.add_argument('--template', required=True, help='模板 .docx 文件路径')
    parser.add_argument('--md', required=True, help='Markdown 内容文件路径')
    parser.add_argument('--output', help='输出 .docx 文件路径')
    parser.add_argument('--mapping', help='章节映射 JSON 文件路径（可选）')
    parser.add_argument('--fig-dir', help='图片目录路径（Markdown 中相对路径的基准目录）')
    args = parser.parse_args()

    if not os.path.exists(args.template):
        print(f"[ERROR] 模板文件不存在: {args.template}")
        sys.exit(1)
    if not os.path.exists(args.md):
        print(f"[ERROR] Markdown 文件不存在: {args.md}")
        sys.exit(1)

    output_path = args.output
    if not output_path:
        base = os.path.splitext(args.md)[0]
        output_path = base + '_from_template.docx'

    fig_dir = args.fig_dir or os.path.dirname(os.path.abspath(args.md))

    mapping = DEFAULT_MAPPING.copy()
    custom_mapping = load_mapping(args.mapping)
    mapping.update(custom_mapping)

    print(f"模板: {args.template}")
    print(f"Markdown: {args.md}")
    print(f"输出: {output_path}")
    print(f"图片目录: {fig_dir}")
    print("---")

    # 解析 Markdown
    sections = parse_markdown(args.md)
    print(f"解析到 {len(sections)} 个章节")
    for s in sections:
        print(f"  {'#' * s['level']} {s['title']} ({len(s['content'])} 段, {len(s['images'])} 图)")

    # 加载模板
    doc = load_template(args.template)

    # 逐章节填充
    for section in sections:
        title = section['title']
        content_lines = [l for l in section['content'] if l.strip()]

        if not content_lines and not section['images']:
            continue

        # 查找锚点
        anchor_text = mapping.get(title)
        anchor_idx = None
        if anchor_text:
            anchor_idx = find_para_index(doc, anchor_text)
        if anchor_idx is None:
            anchor_idx = find_section_anchor(doc, title)

        if anchor_idx is not None:
            print(f"  [OK] '{title}' -> 段落 {anchor_idx} ('{doc.paragraphs[anchor_idx].text[:30]}...')")
        else:
            print(f"  [SKIP] '{title}' -> 未找到模板锚点")

        # 填充内容
        if content_lines:
            fill_section(doc, anchor_idx, content_lines)

        # 填充图片
        if section['images']:
            fill_images(doc, anchor_idx, section['images'], fig_dir)

    save_doc(doc, output_path)
    print(f"\n完成! 已写入: {output_path}")


if __name__ == '__main__':
    main()
